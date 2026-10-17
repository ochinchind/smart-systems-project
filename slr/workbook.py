"""Step 13: assemble the master evidence register (one Excel workbook, 15 sheets).

The CSV files under data/ are the source of truth (diff-able in git); the workbook is a
read-only snapshot regenerated on every run.
"""

from __future__ import annotations

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .project import Project, read_csv

SHEETS = [
    ("Search_Log", lambda p: p.search_log),
    ("Raw_Records", lambda p: p.raw_records),
    ("Deduplication", lambda p: p.dedup_log),
    ("Title_Abstract_Screening", lambda p: p.ta_sheet),
    ("Full_Text_Screening", lambda p: p.ft_sheet),
    ("Included_Studies", lambda p: p.included),
    ("Excluded_Full_Text", None),
    ("Quality_Assessment", lambda p: p.qa_sheet),
    ("Technical_Extraction", lambda p: p.technical),
    ("Quantitative_Extraction", lambda p: p.quantitative),
    ("RQ_Evidence", lambda p: p.rq_evidence),
    ("Research_Gaps", lambda p: p.research_gaps),
    ("PRISMA_Counts", lambda p: p.outputs_dir / "prisma" / "prisma_counts.csv"),
    ("Reference_Audit", lambda p: p.protocol_dir / "reference_audit_template.csv"),
    ("Submission_Audit", lambda p: p.protocol_dir / "submission_audit_template.csv"),
]

LONG_TEXT = {"abstract", "keywords", "affiliations", "raw_ids", "query"}


def build(project: Project) -> str:
    out = project.outputs_dir / "Master_Evidence_Register.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    ft = read_csv(project.ft_sheet)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, getter in SHEETS:
            if name == "Excluded_Full_Text":
                df = ft[ft["final_decision"] == "exclude"] if not ft.empty else pd.DataFrame()
            else:
                df = read_csv(getter(project))
            if df.empty:
                df = pd.DataFrame({"status": ["(no data yet)"]})
            # Excel cell limit is 32 767 characters
            df = df.apply(lambda s: s.astype(str).str.slice(0, 32000))
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for i, col in enumerate(df.columns, start=1):
                ws.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF")
                ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="2A6F97")
                width = 60 if col in LONG_TEXT or col == "title" else min(40, max(10, len(str(col)) + 2))
                ws.column_dimensions[get_column_letter(i)].width = width
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
    return str(out)
