"""Combine every project document into one Word file (and a PDF copy).

    python scripts/build_report.py

Output: outputs/SLR_Project_Documentation.docx and .pdf
Needs pandoc; the PDF additionally needs Chromium/Chrome (set CHROME=/path/to/chrome if not found).
"""

from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from slr.project import read_csv  # noqa: E402
from slr.synthesis import md_table  # noqa: E402

OUT = ROOT / "outputs"
PAGE_BREAK = '\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'

PARTS = [
    ("Part I — Review planning (Steps 1–6)", [
        "docs/phase1_planning/01_Topic_Feasibility_Note.md",
        "docs/phase1_planning/02_Scoping_Search_Report.md",
        "docs/phase1_planning/03_Scope_Statement.md",
        "docs/phase1_planning/04_05_Objectives_and_Research_Questions.md",
        "docs/phase1_planning/06_Review_Protocol.md",
    ]),
    ("Part II — Search strategy (Steps 7–12)", [
        "docs/phase2_search/07_Search_Concept_Matrix.md",
        "docs/phase2_search/08_Database_Selection.md",
        "docs/phase2_search/09_Database_Search_Strings.md",
        "docs/phase2_search/10_Pilot_Search.md",
        "docs/phase2_search/11_Search_Protocol_v1.0_LOCKED.md",
        "docs/phase2_search/12_Multi_Wave_Search_Plan.md",
    ]),
    ("Part III — Record management and PRISMA (Steps 13–17)", [
        "docs/phase3_records/13-17_Record_Management_and_PRISMA.md",
    ]),
    ("Part IV — Quality assessment (Step 18)", [
        "docs/phase4_quality/18_Quality_Assessment.md",
    ]),
    ("Part V — Data extraction (Steps 19–21)", [
        "docs/phase5_extraction/19-21_Data_Extraction.md",
    ]),
    ("Part VI — Classification and synthesis (Steps 22–27)", [
        "docs/phase6_synthesis/22-27_Synthesis_Plan.md",
    ]),
]


LIST_ITEM = re.compile(r"^\s*([*+-]|\d+\.)\s")


def demote(md: str, levels: int = 1) -> str:
    """Shift headings down so each document sits under its Part heading, and add the blank
    line pandoc needs before a list that directly follows a paragraph (GitHub does not need it)."""
    out, in_code = [], False
    for line in md.splitlines():
        if line.startswith("```"):
            in_code = not in_code
        if not in_code:
            if line.startswith("#"):
                line = "#" * levels + line
            prev = out[-1] if out else ""
            if LIST_ITEM.match(line) and prev.strip() and not LIST_ITEM.match(prev) \
                    and not prev.startswith("|") and not prev.startswith("   "):
                out.append("")
        out.append(line)
    return "\n".join(out)


def csv_section(title: str, path: str, cols: list[str] | None = None) -> str:
    df = read_csv(ROOT / path)
    if cols:
        df = df[cols]
    return f"## {title}\n\nSource file: `{path}`\n\n" + md_table(df, max_rows=500) + "\n"


def overview() -> str:
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    body = readme[readme.index("This repository"):]  # title block is on the cover page
    return "# Project overview and status\n\n" + demote(body, 0)


def build_markdown() -> str:
    meta = (
        "---\n"
        'title: "Smart Traffic Lights That Learn"\n'
        'subtitle: "A Systematic Literature Review of Reinforcement Learning for Traffic Signal Control — '
        'From Simulation to Real Streets. Project documentation: Steps 1–27"\n'
        'author: ["Rakhimzhanov Amangeldy", "Nurkin Aidar"]\n'
        f'date: "{date.today().strftime("%d %B %Y")}"\n'
        "---\n\n"
    )
    chunks = [overview()]
    for part, files in PARTS:
        chunks.append(f"# {part}\n")
        for f in files:
            chunks.append(demote((ROOT / f).read_text(encoding="utf-8")))
    chunks.append(
        "# Appendices — protocol data files\n\n"
        + csv_section("Appendix A — Pilot benchmark (gold) set", "data/protocol/known_relevant_papers.csv",
                      ["gold_id", "title", "first_author", "year", "venue", "role"])
        + csv_section("Appendix B — Extraction codebook", "data/protocol/codebook.csv",
                      ["table", "field", "section", "supports", "type", "allowed_values", "description"])
        + csv_section("Appendix C — Methodological families (conceptual comparison)",
                      "data/protocol/method_families.csv")
        + csv_section("Appendix D — Search log (PRISMA-S)", "data/protocol/search_log.csv",
                      ["search_id", "wave", "database", "fields", "filters", "records_retrieved", "run_by"])
    )
    # each top-level part starts on a new page
    body = PAGE_BREAK.join(chunks)
    return meta + body


CSS = """
body { font-family: Calibri, 'Segoe UI', Arial, sans-serif; font-size: 10.5pt; line-height: 1.45;
       color: #1d1d1f; max-width: none; margin: 0; }
h1 { color: #1f4e79; font-size: 20pt; border-bottom: 2px solid #1f4e79; padding-bottom: 4px;
     margin-top: 0; break-before: page; }
h2 { color: #2a6f97; font-size: 14pt; margin-top: 1.4em; }
h3 { color: #2a6f97; font-size: 11.5pt; }
table { border-collapse: collapse; width: 100%; margin: 0.6em 0 1em; font-size: 8.5pt; }
th { background: #d9e7f2; text-align: left; }
th, td { border: 1px solid #a6a6a6; padding: 3px 5px; vertical-align: top; }
tr { break-inside: avoid; }
code { font-family: Consolas, 'DejaVu Sans Mono', monospace; font-size: 8.5pt; }
pre { background: #f4f6f8; border: 1px solid #dde3e8; padding: 6px 8px; white-space: pre-wrap;
      word-break: break-word; font-size: 8pt; }
blockquote { border-left: 3px solid #2a6f97; margin-left: 0; padding-left: 10px; color: #444; }
header#title-block-header { text-align: center; padding-top: 30vh; break-after: page; }
header#title-block-header .title { font-size: 28pt; color: #1f4e79; border: none; break-before: avoid; }
header#title-block-header .subtitle { font-size: 14pt; color: #2a6f97; margin: 0 10% 2em; }
header#title-block-header .author, header#title-block-header .date { font-size: 12pt; margin: 0.2em; }
nav#TOC { break-after: page; }
nav#TOC::before { content: "Contents"; display: block; font-size: 20pt; color: #1f4e79; font-weight: bold;
                  border-bottom: 2px solid #1f4e79; margin-bottom: 0.6em; }
nav#TOC ul { list-style: none; padding-left: 1.2em; }
nav#TOC > ul { padding-left: 0; }
nav#TOC > ul > li > a { font-weight: bold; }
nav#TOC a { color: #1d1d1f; text-decoration: none; }
a { color: #2a6f97; }
@page { size: A4; margin: 18mm 16mm 18mm 16mm; }
"""


def find_chrome() -> str | None:
    for c in (os.environ.get("CHROME"), "/opt/pw-browsers/chromium-1194/chrome-linux/chrome",
              shutil.which("chromium"), shutil.which("chromium-browser"), shutil.which("google-chrome")):
        if c and Path(c).exists():
            return c
    found = sorted(Path("/opt/pw-browsers").glob("chromium-*/chrome-linux/chrome"))
    return str(found[-1]) if found else None


def build_pdf(md: Path, pdf: Path) -> bool:
    chrome = find_chrome()
    if not chrome:
        return False
    css = md.with_suffix(".css")
    css.write_text(CSS, encoding="utf-8")
    html = md.with_suffix(".html")
    subprocess.run(["pandoc", str(md), "-o", str(html), "--standalone", "--toc", "--toc-depth=2",
                    "--embed-resources", "--css", str(css), "--metadata", "lang=en",
                    "--from", "markdown+pipe_tables+raw_attribute"], check=True)
    subprocess.run([chrome, "--headless", "--no-sandbox", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={pdf}", html.as_uri()], check=True, capture_output=True)
    css.unlink()
    html.unlink()
    return True


def style_docx(path: Path) -> None:
    doc = Document(path)
    # ask Word to refresh the table of contents when the file is opened
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in doc.styles]:
            doc.styles[name].font.name = "Calibri"
            doc.styles[name].font.size = Pt(10.5)
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "999999")
            borders.append(el)
        tbl_pr.append(borders)
        for cell in table.rows[0].cells:
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "D9E7F2")
            cell._tc.get_or_add_tcPr().append(shd)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)
    doc.save(path)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    md = OUT / "SLR_Project_Documentation.md"
    md.write_text(build_markdown(), encoding="utf-8")
    docx = OUT / "SLR_Project_Documentation.docx"
    subprocess.run(["pandoc", str(md), "-o", str(docx), "--toc", "--toc-depth=2",
                    "--from", "markdown+pipe_tables+raw_attribute"], check=True)
    style_docx(docx)
    print("written", docx)
    pdf = OUT / "SLR_Project_Documentation.pdf"
    if build_pdf(md, pdf):
        print("written", pdf)
    else:
        print("Chromium/Chrome not found - PDF skipped (open the .docx and use Save as PDF)")
    md.unlink()


if __name__ == "__main__":
    main()
